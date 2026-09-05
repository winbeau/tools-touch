"""Reusable 985/211 deadline report from public data or a saved snapshot."""
import argparse
import csv
from copy import copy
from collections import Counter
from datetime import date, datetime
import json
from pathlib import Path
import re
import tempfile

from baoyan import Client, row, safe_cell
from baoyan_subjects import add_subject_labels
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pypinyin import lazy_pinyin


def tiers(value):
    return {t.strip().removesuffix("工程") for t in re.split(r"[,，/、]", value or "")}.intersection({"985", "211"})


# School identity is independent of the site's occasionally incorrect tier tags.
SCHOOLS_985 = set("北京大学 清华大学 中国人民大学 北京师范大学 北京航空航天大学 北京理工大学 中国农业大学 中央民族大学 南开大学 天津大学 大连理工大学 东北大学 吉林大学 哈尔滨工业大学 复旦大学 上海交通大学 同济大学 华东师范大学 南京大学 东南大学 浙江大学 中国科学技术大学 厦门大学 山东大学 中国海洋大学 武汉大学 华中科技大学 湖南大学 中南大学 国防科技大学 中山大学 华南理工大学 四川大学 电子科技大学 重庆大学 西安交通大学 西北工业大学 西北农林科技大学 兰州大学".split())
SCHOOLS_211 = SCHOOLS_985 | set("北京交通大学 北京工业大学 北京科技大学 北京化工大学 北京邮电大学 北京林业大学 北京中医药大学 北京外国语大学 中国传媒大学 中央财经大学 对外经济贸易大学 北京体育大学 中央音乐学院 中国政法大学 华北电力大学 天津医科大学 河北工业大学 太原理工大学 内蒙古大学 辽宁大学 大连海事大学 延边大学 东北师范大学 哈尔滨工程大学 东北农业大学 东北林业大学 华东理工大学 东华大学 上海外国语大学 上海财经大学 上海大学 海军军医大学 苏州大学 南京航空航天大学 南京理工大学 中国矿业大学 河海大学 江南大学 南京农业大学 中国药科大学 南京师范大学 安徽大学 合肥工业大学 福州大学 南昌大学 中国石油大学（华东） 中国石油大学（北京） 郑州大学 中国地质大学（武汉） 中国地质大学（北京） 武汉理工大学 华中农业大学 华中师范大学 中南财经政法大学 湖南师范大学 暨南大学 华南师范大学 广西大学 海南大学 西南交通大学 四川农业大学 西南财经大学 西南大学 贵州大学 云南大学 西藏大学 西北大学 西安电子科技大学 长安大学 陕西师范大学 空军军医大学 青海大学 宁夏大学 新疆大学 石河子大学 中国矿业大学（北京）".split())
ALIASES = {"哈尔滨工业大学（深圳）": "哈尔滨工业大学", "哈尔滨工业大学（威海）": "哈尔滨工业大学", "华北电力大学（北京）": "华北电力大学", "华北电力大学（保定）": "华北电力大学", "中国矿业大学（徐州）": "中国矿业大学", "国防科学技术大学": "国防科技大学", "第二军医大学": "海军军医大学", "第四军医大学": "空军军医大学"}


def university(name):
    return ALIASES.get(name or "", name or "")


def school_key(name):
    original = name or ""
    name = university(name)
    reading = {"重庆大学": "chongqingdaxue", "厦门大学": "xiamendaxue", "长安大学": "changandaxue"}
    rank = 0 if name in SCHOOLS_985 else 1 if name in SCHOOLS_211 else 2
    return rank, reading.get(name) or "".join(lazy_pinyin(name)).lower(), "".join(lazy_pinyin(original)).lower(), original


def activity_end_key(value):
    try:
        parsed = datetime.fromisoformat(str(value).strip().replace("/", "-"))
        return -(parsed.toordinal() * 86400 + parsed.hour * 3600 + parsed.minute * 60 + parsed.second)
    except ValueError:
        return float("inf")


def sort_rows(rows):
    rows.sort(key=lambda r: (school_key(r.get("学校")), activity_end_key(r.get("活动结束时间")),
                             r.get("学校") or "", r.get("学院") or "", r.get("信息ID") or 0))


def format_table(ws, headers):
    """Fixed widths, single-line cells and visible borders over school colors."""
    widths = {"学校": 32, "学院": 40, "标题": 65, "数据说明": 60, "筛选依据": 60, "说明": 65,
              "信息ID": 12, "院校层次": 18, "筛选建议": 14, "学科筛选标签": 24}
    for index, header in enumerate(headers, 1):
        dimension = ws.column_dimensions[get_column_letter(index)]
        dimension.width = widths.get(header, 55 if "链接" in header else 25)
        dimension.bestFit = False
    edge = Side(style="thin", color="B8C4D0")
    ws.sheet_view.showGridLines = True
    for cells in ws.iter_rows():
        for cell in cells:
            if isinstance(cell.value, str):
                cell.value = re.sub(r"[\r\n\t]+", " ", cell.value)
            cell.alignment = Alignment(vertical="center", wrap_text=False, shrink_to_fit=True)
            # Preserve any existing explicit edges; add a thin border where absent.
            border = copy(cell.border)
            for side in ("left", "right", "top", "bottom"):
                existing = getattr(border, side)
                if existing is None or existing.style is None:
                    setattr(border, side, edge)
            cell.border = border
            if headers[cell.column - 1] in ("学校", "活动结束时间"):
                font = copy(cell.font)
                font.bold = True
                cell.font = font
        ws.row_dimensions[cells[0].row].height = 24 if cells[0].row == 1 else 22


def add_arguments(parser):
    parser.add_argument("--year", type=int, default=2026, help="网站年份，默认2026")
    parser.add_argument("--deadline-from", type=date.fromisoformat, default=date(2026, 9, 5), metavar="YYYY-MM-DD", help="报名截止日期下限（含当天），默认2026-09-05")
    parser.add_argument("--snapshot", type=Path, help="复用含原始保研信息、学校目录和查询说明JSON的导出目录，不联网")
    parser.add_argument("--output-dir", type=Path, help="新导出目录；已存在时拒绝覆盖，默认自动创建时间戳目录")
    parser.set_defaults(func=run)


def run(args):
    if args.deadline_from.year != args.year:
        raise ValueError("--deadline-from 的年份必须与 --year 一致")
    destination = args.output_dir or Path(__file__).parent / "exports" / (f"985-211-预推免-{args.year}-分组排序-" + datetime.now().strftime("%Y%m%d-%H%M%S"))
    if destination.exists():
        raise RuntimeError("导出目录已存在，请使用新的 --output-dir；不会覆盖既有文件")
    destination.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=".baoyan-report-", dir=destination.parent) as staging:
        build_report(args, Path(staging))
        Path(staging).rename(destination)
    print("OUTPUT=" + str(destination.resolve()), flush=True)


def build_report(args, destination):
    captured = datetime.now().astimezone().isoformat(timespec="seconds")
    snapshot = args.snapshot
    client = Client(anonymous=True) if not snapshot else None
    schools = json.loads((snapshot / "学校目录.json").read_text()) if snapshot else []
    school_ids = set()
    school_total = None
    for page in ([] if snapshot else range(1, 101)):
        response = client.get("/search/colleges", page=page, size=100)
        school_total = response["total_count"]
        if not response["content"]:
            break
        for school in response["content"]:
            if school["id"] in school_ids:
                raise RuntimeError("学校目录出现重复ID")
            schools.append(school)
            school_ids.add(school["id"])
        if len(schools) == school_total:
            break
    if snapshot:
        previous_metadata = json.loads((snapshot / "查询说明.json").read_text())
        school_total = previous_metadata["目录声明学校数"]
        captured = previous_metadata["抓取时间"]
    target_schools = {s["name"]: s for s in schools if ALIASES.get(s["name"], s["name"]) in SCHOOLS_211}
    print(f"学校目录实际返回 {len(schools)} / 声明 {school_total} 条；按独立学校名称名单筛选全量通知。", flush=True)
    records = []
    source = json.loads((snapshot / "全部原始保研信息.json").read_text()) if snapshot else client.pages("/articles", category="保研信息", all=1)
    for item in source:
        records.append(item)
        if len(records) % 1000 == 0:
            print(f"已完整读取 {len(records)} 条保研信息…", flush=True)
    if len({item["id"] for item in records}) != len(records):
        raise RuntimeError("原始通知包含重复ID")
    if snapshot and len(records) != previous_metadata["扫描保研信息数"]:
        raise RuntimeError("快照记录数与查询说明不一致")
    print(f"{'快照数量与ID校验' if snapshot else '全部分页校验'}完成：{len(records)} 条，无重复ID", flush=True)
    selected, uncertain = [], []
    counts = Counter()
    current_by_school = Counter()
    for item in records:
        if "预推免" not in {t.strip() for t in (item.get("tags") or "").split(",")}:
            continue
        name = item.get("college") or ""
        canonical = ALIASES.get(name, name)
        if canonical not in SCHOOLS_211:
            if tiers(item.get("college_level")) and str(item.get("sign_up_end") or "").startswith(str(args.year)) and str(item.get("sign_up_end")) >= args.deadline_from.isoformat():
                uncertain.append({"院校层次": "待核实", **row(item), "待核实原因": "学校名称缺失或网站985/211标签与实际院校身份不符，未纳入匹配"})
                counts["school_identity_conflict"] += 1
            continue
        level = "985/211" if canonical in SCHOOLS_985 else "211（非985）"
        result = {"院校层次": level, **row(item)}
        result["数据说明"] = "；".join(filter(None, [result["数据说明"], "保研信息网字段；未核验官网或登录报名系统；截止日不代表目前仍可报名"]))
        if not item.get("year"):
            if str(item.get("sign_up_end") or "").startswith(str(args.year)) and str(item.get("sign_up_end")) >= args.deadline_from.isoformat():
                result["待核实原因"] = "截止日期符合条件，但网站年份缺失；单独列出，不推断网站年份"
                uncertain.append(result)
                counts["unknown_year_matching_deadline"] += 1
            continue
        if str(item["year"]) != str(args.year):
            continue
        current_by_school[item["college"]] += 1
        end = item.get("sign_up_end") or ""
        try:
            deadline = date.fromisoformat(end[:10].replace("/", "-"))
        except ValueError:
            result["待核实原因"] = "报名截止日期缺失或无法解析"
            uncertain.append(result)
            counts["unknown_deadline"] += 1
            continue
        if deadline.year != args.year:
            result["待核实原因"] = f"报名截止年份与网站年份{args.year}不一致"
            uncertain.append(result)
            counts["year_conflict"] += 1
            continue
        if deadline >= args.deadline_from:
            selected.append(result)
        else:
            counts["before_threshold"] += 1
    sort_rows(selected)
    sort_rows(uncertain)
    selected = [add_subject_labels(r) for r in selected]
    uncertain = [add_subject_labels(r) for r in uncertain]
    preferred = [r for r in selected if r["筛选建议"] == "保留"]
    subject_uncertain = [r for r in selected if r["筛选建议"] == "待核实"]
    excluded = [r for r in selected if r["筛选建议"] == "排除"]
    grouped = {}
    for r in selected:
        grouped.setdefault(university(r["学校"]), []).append(r)
    summary = [{"学校": name, "院校层次": rows[0]["院校层次"], "符合条件通知数": len(rows),
                "学科优先保留数": sum(r["筛选建议"] == "保留" for r in rows),
                "学科待核实数": sum(r["筛选建议"] == "待核实" for r in rows),
                "学科建议排除数": sum(r["筛选建议"] == "排除" for r in rows),
                "最早报名截止": min(r["报名截止时间"] for r in rows),
                "最晚报名截止": max(r["报名截止时间"] for r in rows),
                "原始学校名称／校区": "、".join(sorted({r["学校"] for r in rows}))} for name, rows in grouped.items()]
    summary.sort(key=lambda r: school_key(r["学校"]))
    coverage = [{"学校": name, "院校层次": "985/211" if name in SCHOOLS_985 else "211（非985）",
                 f"网站{args.year}预推免通知数": sum(n for school, n in current_by_school.items() if ALIASES.get(school, school) == name),
                 "符合条件通知数": sum(len(rows) for school, rows in grouped.items() if ALIASES.get(school, school) == name),
                 "说明": "同一学校校区合并统计；无匹配不代表今年已截止或没有招生"}
                for name in sorted(SCHOOLS_211, key=school_key)]
    metadata = {"抓取时间": captured, "来源": "http://pc.baoyanwang.com.cn/", "API": "http://api.baoyanwang.com.cn/api/v1",
                "查询方式": "本地快照重导出，未联网" if snapshot else "baoyan.Client(anonymous=True)，只读公开接口，全部分页后本地筛选",
                "筛选条件": f"985/211学校名称名单（校区别名归属母校）；标签预推免；网站year={args.year}；报名截止日期>={args.deadline_from}且截止年份={args.year}（含当天）；年份空白但截止日期符合的记录另列待核实",
                "排序与底色": "985在前，非985的211在后；组内大学名称按全拼字母升序，同大学原始名称连续，校区相邻且与母校同色；每个原始学校名称内按活动结束时间降序，缺失或无效时间置后；未使用商业大学排名",
                "学科筛选规则": "名称规则初筛：保留工科、理科和教育类（教育例外优先）；排除经管财经、文史社科外语、医药农林体艺；医农工程交叉保留，文理交叉及全校通知另列待核实。不按财经/外国语大学校名一刀切，不改变原始日期匹配结果。",
                "学科筛选统计": dict(Counter(r["筛选建议"] for r in selected)),
                "学科标签统计": dict(Counter(r["学科筛选标签"] for r in selected)),
                "排版": "固定列宽和行高，关闭自动换行，长内容缩小字体以在单行内显示；保留已有边框，空白边补细边框；学校与活动结束时间列加粗；Excel中的原始换行与制表符仅显示为空格，JSON保留原文",
                "完整性": f"全部保研信息原抓取经Client.pages校验total_count及重复ID；全部仅指网站现存记录。学校目录声明{school_total}，实际返回{len(schools)}；使用查询脚本中的985/211学校名称名单匹配全量通知，不依赖目录完整性",
                "限制": "未访问官网核验；报名截止时间使用sign_up_end，非活动结束时间；学校汇总范围不是统一截止日期；未知年份及截止日期另列，不推断为已结束；网站可能默认填入23:59，未核验是否精确到分钟",
                "目录声明学校数": school_total, "目录实际返回学校数": len(schools), "目录985及211学校数": len(target_schools), "筛选名单学校数": len(SCHOOLS_211), "扫描保研信息数": len(records),
                "匹配通知数": len(selected), "匹配学校数": len(grouped), "待核实通知数": len(uncertain), "排除与异常计数": dict(counts)}
    # Output data files are generated artifacts, not changes to application code.
    for name, data in [("全部原始保研信息", records), ("学校目录", schools), ("匹配明细", selected), ("优先保留", preferred), ("学科待核实", subject_uncertain), ("建议排除", excluded), ("查询说明", metadata)]:
        (destination / f"{name}.json").write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    wb = Workbook()
    wb.remove(wb.active)
    palette = ["EAF3FF", "FFF2DB", "E8F5E9", "F3E9FA", "E0F3F3", "FCE8EC"]
    visible = list(dict.fromkeys(university(r["学校"]) for r in selected))
    remaining = {university(r["学校"]) for r in uncertain + coverage} - set(visible)
    ordered_schools = visible + sorted(remaining, key=school_key)
    school_colors = {name: palette[index % len(palette)] for index, name in enumerate(ordered_schools)}
    tables = [("优先保留", preferred), ("匹配明细", selected), ("学科待核实", subject_uncertain), ("建议排除", excluded),
              ("学校汇总", summary), ("待核实", uncertain), ("院校覆盖", coverage),
              ("查询说明", [{"项目": k, "说明": json.dumps(v, ensure_ascii=False) if isinstance(v, dict) else v} for k, v in metadata.items()])]
    for title, rows in tables:
        ws = wb.create_sheet(title)
        if not rows:
            ws.append(["无记录"])
            continue
        headers = list(rows[0])
        ws.append(headers)
        for r in rows:
            ws.append([safe_cell(r.get(h, "")) for h in headers])
        ws.freeze_panes = "D2" if "学科筛选标签" in headers else "A2"
        ws.auto_filter.ref = ws.dimensions
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="24598A")
        for data, cells in zip(rows, ws.iter_rows(min_row=2)):
            for cell in cells:
                if "链接" in headers[cell.column-1] and isinstance(cell.value, str) and cell.value.startswith(("http://", "https://")):
                    cell.hyperlink = cell.value
                    cell.style = "Hyperlink"
                if "学校" in data:
                    cell.fill = PatternFill("solid", fgColor=school_colors[university(data["学校"])])
        format_table(ws, headers)
        with (destination / f"{title}.csv").open("w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=headers)
            writer.writeheader()
            writer.writerows({k: safe_cell(r.get(k, "")) for k in headers} for r in rows)
    wb.save(destination / "985-211预推免截止日期筛选.xlsx")
    print(json.dumps(metadata, ensure_ascii=False, indent=2), flush=True)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    add_arguments(parser)
    try:
        run(parser.parse_args())
    except (RuntimeError, ValueError, OSError, KeyError) as error:
        parser.exit(1, f"错误：{error}\n")
