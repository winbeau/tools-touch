import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from baoyan import Client, FIELDS, write_table


class PaginationTests(unittest.TestCase):
    def client(self, pages):
        c = Client.__new__(Client)
        c.get = lambda path, **params: pages[params['page'] - 1]
        return c

    def test_server_caps_page_size(self):
        c = self.client([{'content': [{'id': 1}], 'total_count': 2},
                         {'content': [{'id': 2}], 'total_count': 2}])
        self.assertEqual([x['id'] for x in c.pages('/articles', size=100)], [1, 2])

    def test_repeated_page_fails(self):
        c = self.client([{'content': [{'id': 1}], 'total_count': 2}] * 2)
        with self.assertRaisesRegex(RuntimeError, '重复'):
            list(c.pages('/articles'))

    def test_missing_page_fails(self):
        c = self.client([{'content': [{'id': 1}], 'total_count': 2},
                         {'content': [], 'total_count': 2}])
        with self.assertRaisesRegex(RuntimeError, '不完整'):
            list(c.pages('/articles'))

    def test_changed_total_fails(self):
        c = self.client([{'content': [{'id': 1}], 'total_count': 2},
                         {'content': [{'id': 2}], 'total_count': 3}])
        with self.assertRaisesRegex(RuntimeError, '变化'):
            list(c.pages('/articles'))

    def test_excel_formula_escaped_and_link_preserved(self):
        r = {label: '' for _, label in FIELDS}
        r.update({'标题': '=1+1', '官网链接': 'https://example.edu.cn/a'})
        with tempfile.TemporaryDirectory() as d:
            dest = Path(d) / 'test.xlsx'
            write_table(dest, [r])
            ws = load_workbook(dest).active
            self.assertEqual(ws['D2'].data_type, 's')
            self.assertEqual(ws['K2'].hyperlink.target, r['官网链接'])


if __name__ == '__main__':
    unittest.main()
