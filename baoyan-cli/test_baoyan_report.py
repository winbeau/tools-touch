import argparse
from contextlib import redirect_stdout
from datetime import date
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from baoyan_report import run, school_key, sort_rows
from baoyan_subjects import classify_subject


class ReportTests(unittest.TestCase):
    def test_subject_preferences_and_education_exception(self):
        examples = [
            ("经济管理学院", "排除"), ("商学院", "排除"), ("外国语学院", "排除"),
            ("金融工程研究中心", "排除"), ("国际中文教育学院", "保留"),
            ("教育学部", "保留"), ("生物医学工程学院", "保留"),
            ("计算机学院", "保留"), ("物理学院", "保留"),
            ("信息管理学院", "待核实"), ("建筑与艺术学院", "待核实"),
            ("全校通知", "待核实"), ("药学院", "排除"),
        ]
        for academy, expected in examples:
            with self.subTest(academy=academy):
                self.assertEqual(classify_subject({"学院": academy})["筛选建议"], expected)
        self.assertEqual(classify_subject({"学院": "", "标题": "【中央财经大学】——计算机学院"})["筛选建议"], "保留")
        self.assertEqual(classify_subject({"学院": "", "标题": "【北京外国语大学】——全校通知"})["筛选建议"], "待核实")
        self.assertEqual(classify_subject({"学院": "", "标题": "【北京大学】——教育部重点实验室（材料工程）"})["学科筛选标签"], "工科")

    def test_pinyin_and_tier_order(self):
        names = ["厦门大学", "北京交通大学", "重庆大学", "北京大学", "安徽大学", "长安大学"]
        self.assertEqual(sorted(names, key=school_key),
                         ["北京大学", "重庆大学", "厦门大学", "安徽大学", "北京交通大学", "长安大学"])

    def test_school_groups_use_activity_end_not_signup_deadline(self):
        rows = [
            {"信息ID": 1, "学校": "北京大学", "活动结束时间": "2026-09-10 12:00", "报名截止时间": "2026-09-20"},
            {"信息ID": 2, "学校": "安徽大学", "活动结束时间": "2026-12-01"},
            {"信息ID": 3, "学校": "北京大学", "活动结束时间": "2026/09/12 09:00", "报名截止时间": "2026-09-05"},
            {"信息ID": 4, "学校": "北京大学", "活动结束时间": "待通知"},
            {"信息ID": 5, "学校": "北京大学", "活动结束时间": None},
        ]
        sort_rows(rows)
        self.assertEqual([r["信息ID"] for r in rows], [3, 1, 4, 5, 2])

    def test_snapshot_export_preserves_ids_and_applies_group_colors(self):
        records = []
        for identifier, school, end in [(1, "北京大学", "2026-09-11"), (2, "安徽大学", "2026-10-01"),
                                         (3, "北京大学", "2026-09-15"), (4, "哈尔滨工业大学（深圳）", "2026-09-20"),
                                         (5, "哈尔滨工业大学", "2026-09-10")]:
            academy = {1: "经济管理学院", 2: "教育学院", 3: "计算机学院", 4: "外国语学院", 5: "交叉研究院"}[identifier]
            records.append(dict(id=identifier, college=school, academy=academy, tags="预推免", year=2026,
                                sign_up_end="2026-09-05 23:59", end_time=end,
                                office_url="https://example.edu.cn/notice", sign_up_url="https://example.edu.cn/signup"))
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            snapshot = root / "snapshot"
            snapshot.mkdir()
            for name, data in [("全部原始保研信息", records), ("学校目录", []),
                               ("查询说明", {"扫描保研信息数": 5, "目录声明学校数": 0, "抓取时间": "fixture"})]:
                (snapshot / (name + ".json")).write_text(json.dumps(data), encoding="utf-8")
            args = argparse.Namespace(year=2026, deadline_from=date(2026, 9, 5), snapshot=snapshot, output_dir=root / "output")
            with patch("baoyan_report.Client", side_effect=AssertionError("snapshot must not access network")), redirect_stdout(io.StringIO()):
                run(args)
            rows = json.loads((args.output_dir / "匹配明细.json").read_text(encoding="utf-8"))
            self.assertEqual([r["信息ID"] for r in rows], [3, 1, 5, 4, 2])
            wb = load_workbook(args.output_dir / "985-211预推免截止日期筛选.xlsx")
            ws = wb["匹配明细"]
            self.assertEqual(ws.max_row, 6)
            self.assertEqual(ws["A2"].fill.fgColor.rgb, ws["A3"].fill.fgColor.rgb)
            self.assertEqual(ws["A4"].fill.fgColor.rgb, ws["A5"].fill.fgColor.rgb)
            self.assertNotEqual(ws["A3"].fill.fgColor.rgb, ws["A4"].fill.fgColor.rgb)
            for cells in ws.iter_rows(min_row=2):
                self.assertTrue(all(c.fill.fgColor.rgb == cells[0].fill.fgColor.rgb for c in cells))
            headers = [c.value for c in ws[1]]
            self.assertEqual(ws.cell(2, headers.index("官网链接")+1).hyperlink.target, records[0]["office_url"])
            self.assertEqual(ws.freeze_panes, "D2")
            self.assertEqual(wb["学校汇总"].max_row, 4)
            self.assertEqual(wb.sheetnames[0], "优先保留")
            self.assertEqual(wb["优先保留"].max_row, 3)
            self.assertEqual(wb["建议排除"].max_row, 3)
            self.assertEqual(wb["学科待核实"].max_row, 2)
            for heading in ("学校", "活动结束时间"):
                self.assertTrue(all(ws.cell(n, headers.index(heading)+1).font.bold for n in range(1, ws.max_row+1)))
            self.assertNotEqual(ws.column_dimensions["A"].width, ws.column_dimensions["C"].width)
            self.assertFalse(ws["C2"].alignment.wrap_text)
            self.assertTrue(ws["C2"].alignment.shrink_to_fit)
            self.assertEqual(ws.column_dimensions["C"].width, 32)
            self.assertTrue(all(not d.bestFit for d in ws.column_dimensions.values()))
            for cells in ws.iter_rows():
                self.assertTrue(all(getattr(c.border, side).style == "thin" for c in cells for side in ("left", "right", "top", "bottom")))
            wb.close()
            with self.assertRaisesRegex(RuntimeError, "已存在"):
                run(args)


if __name__ == "__main__":
    unittest.main()
