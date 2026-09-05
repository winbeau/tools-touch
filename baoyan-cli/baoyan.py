#!/usr/bin/env python3
"""保研信息网 CLI：扫码登录、按学校导出 CSV/XLSX/JSON。"""
import argparse
import csv
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import sys
import tempfile
import time
import uuid

import httpx

SITE = 'http://pc.baoyanwang.com.cn'
API = 'http://api.baoyanwang.com.cn/api/v1'
SESSION = Path.home() / '.config/baoyan-cli/session.json'
PROFILE = Path.home() / '.local/share/baoyan-cli/chrome'
# 网站公开前端使用的请求签名常量，不是用户的登录凭据。
SIGN_SECRET = 'e3fa66d113ae5dd8f291e10209e57bcf'
FIELDS = [('id', '信息ID'), ('college', '学校'), ('academy', '学院'),
          ('title', '标题'), ('tags', '活动类型'), ('year', '网站年份'),
          ('sign_up_start', '报名开始时间'), ('sign_up_end', '报名截止时间'),
          ('start_time', '活动开始时间'), ('end_time', '活动结束时间'),
          ('office_url', '官网链接'), ('sign_up_url', '报名链接'),
          ('sign_up_email', '报名邮箱'), ('updated_time', '更新日期'),
          ('site_url', '信息页链接'), ('notes', '数据说明')]


class Client:
    def __init__(self, anonymous=False, delay=0.3):
        self.token = ''
        if not anonymous:
            self.token = os.environ.get('BAOYAN_TOKEN', '')
            if not self.token and SESSION.exists():
                self.token = json.loads(SESSION.read_text())['token']
        self.delay = delay
        self.last = 0.0
        self.http = httpx.Client(base_url=API, timeout=30, headers={'Referer': SITE + '/', 'User-Agent': 'baoyan-cli/1.0'})

    def get(self, path, **params):
        for attempt in range(3):
            time.sleep(max(0, self.delay - (time.monotonic() - self.last)))
            nonce = f'{uuid.uuid4()}-{int(time.time() * 1000)}'
            headers = {'X-Auth-Nonce': nonce, 'X-Auth-Device': 'web',
                       'X-Auth-Sign': hashlib.md5(f'nonce={nonce}&secret={SIGN_SECRET}'.encode()).hexdigest()}
            if self.token:
                headers['X-Auth-Key'] = self.token
            self.last = time.monotonic()
            try:
                response = self.http.get(path, params=params, headers=headers)
                if response.status_code == 429 or response.status_code >= 500:
                    if attempt < 2:
                        time.sleep(2 ** (attempt + 1))
                        continue
                response.raise_for_status()
            except (httpx.TimeoutException, httpx.NetworkError):
                if attempt < 2:
                    time.sleep(2 ** (attempt + 1))
                    continue
                raise RuntimeError('网络请求失败，请稍后重试。') from None
            data = response.json()
            if data.get('code') == 401:
                raise RuntimeError('登录已失效，请运行 ./baoyan login 重新扫码。')
            if not data.get('success'):
                raise RuntimeError(f"接口拒绝请求：{data.get('errors') or data.get('code')}。可运行 ./baoyan login 后重试。")
            if data.get('encrypt'):
                raise RuntimeError('列表接口变更为加密返回，需要更新脚本。')
            return data['result']

    def pages(self, path, size=100, **params):
        seen = set()
        expected = None
        for page in range(1, 10001):
            result = self.get(path, page=page, size=size, **params)
            if not isinstance(result, dict) or not isinstance(result.get('content'), list):
                raise RuntimeError('列表结构发生变化，已停止导出。')
            total = int(result['total_count'])
            if expected is None:
                expected = total
            elif total != expected:
                raise RuntimeError('抓取期间总条数发生变化，请重新运行以保证完整性。')
            batch = result['content']
            for item in batch:
                if item['id'] in seen:
                    raise RuntimeError('分页出现重复记录，请重新运行；不会输出不完整文件。')
                seen.add(item['id'])
                yield item
            if len(seen) == total:
                return
            if not batch or len(seen) > total:
                raise RuntimeError(f'分页不完整：收到 {len(seen)} / {total} 条。')
        raise RuntimeError('分页超过安全上限。')


def save_token(token):
    SESSION.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
    fd = os.open(SESSION, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
    os.fchmod(fd, 0o600)
    with os.fdopen(fd, 'w') as f:
        json.dump({'token': token}, f)


def login(args):
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        if args.cdp:
            browser = p.chromium.connect_over_cdp(args.cdp)
            context = browser.contexts[0]
        else:
            PROFILE.mkdir(parents=True, exist_ok=True, mode=0o700)
            context = p.chromium.launch_persistent_context(str(PROFILE), channel='chrome', headless=False,
                                                        no_viewport=True)
        page = next((x for x in context.pages if 'pc.baoyanwang.com.cn' in x.url), None)
        if page is None:
            page = context.new_page()
        page.goto(SITE + '/sign/in', wait_until='domcontentloaded')
        print('请在 Chrome 中微信扫码；登录成功后自动保存凭据。', flush=True)
        deadline = time.monotonic() + args.timeout
        while time.monotonic() < deadline:
            token = next((x['value'] for x in context.cookies(SITE) if x['name'] == 'token'), '')
            if token:
                save_token(token)
                print(f'登录已保存：{SESSION}（仅当前用户可读写）')
                if not args.cdp:
                    context.close()
                return
            page.wait_for_timeout(1000)
        raise RuntimeError('等待扫码超时，请重新运行 login。')


def row(item):
    result = {label: item.get(key, '') for key, label in FIELDS}
    result['信息页链接'] = f"{SITE}/articles/{item['id']}"
    notes = []
    if not item.get('year'):
        notes.append('网站年份未填写；未从日期推断')
    for key, label in FIELDS:
        if key in ('sign_up_start', 'sign_up_end', 'office_url') and not item.get(key):
            notes.append(label + '未提供')
    result['数据说明'] = '；'.join(notes)
    return result


def safe_cell(value):
    if value is None:
        return ''
    if isinstance(value, (list, dict)):
        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, str):
        value = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', value)
        if value.lstrip().startswith(('=', '+', '-', '@')):
            value = "'" + value
    return value


def write_table(destination, rows):
    suffix = destination.suffix.lower()
    if suffix not in ('.csv', '.xlsx', '.json'):
        raise RuntimeError('输出扩展名必须是 .csv、.xlsx 或 .json。')
    destination.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(dir=destination.parent, suffix=suffix)
    os.close(fd)
    try:
        headers = [label for _, label in FIELDS]
        if suffix == '.json':
            Path(tmp).write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding='utf-8')
        elif suffix == '.csv':
            with open(tmp, 'w', newline='', encoding='utf-8-sig') as f:
                w = csv.DictWriter(f, fieldnames=headers)
                w.writeheader()
                w.writerows({k: safe_cell(v) for k, v in r.items()} for r in rows)
        else:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
            from openpyxl.utils import get_column_letter
            wb = Workbook()
            ws = wb.active
            ws.title = '保研信息'
            ws.append(headers)
            for r in rows:
                ws.append([safe_cell(r[h]) for h in headers])
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='24598A')
            for i, h in enumerate(headers, 1):
                ws.column_dimensions[get_column_letter(i)].width = 55 if h in ('标题', '数据说明') or '链接' in h else 23
            for line in ws.iter_rows(min_row=2):
                for cell in line:
                    if '链接' in headers[cell.column - 1] and isinstance(cell.value, str) and cell.value.startswith(('http://', 'https://')):
                        cell.hyperlink = cell.value
                        cell.style = 'Hyperlink'
            wb.save(tmp)
        os.replace(tmp, destination)
    finally:
        Path(tmp).unlink(missing_ok=True)


def export(args):
    destination = Path(args.output)
    if destination.exists() and not args.force:
        raise RuntimeError('输出文件已存在；请换文件名或加 --force。')
    if destination.suffix.lower() not in ('.csv', '.xlsx', '.json'):
        raise RuntimeError('输出扩展名必须是 .csv、.xlsx 或 .json。')
    c = Client(args.anonymous, args.delay)
    schools = list(c.pages('/search/colleges', name=args.school))
    exact = [x for x in schools if x['name'] == args.school]
    if not exact:
        suggestions = '、'.join(x['name'] for x in schools) or '无匹配学校'
        raise RuntimeError(f'请使用学校完整名称。搜索结果：{suggestions}')
    records = list(c.pages('/articles', size=args.page_size, college=args.school, category='保研信息', all=1))
    if any(x.get('college') != args.school for x in records):
        raise RuntimeError('接口未正确按学校筛选，已停止导出。')
    kinds = {'夏令营', '预推免'} if args.kind == '全部' else {args.kind}
    selected = [x for x in records if kinds.intersection((x.get('tags') or '').split(','))]
    unknown = sum(not x.get('year') for x in selected)
    if args.year:
        selected = [x for x in selected if str(x.get('year')) == str(args.year)]
        if unknown:
            print(f'提示：{unknown} 条记录的网站年份未填写，已被 --year 排除；不传 --year 可保留。', file=sys.stderr)
    selected.sort(key=lambda x: (x.get('sign_up_end') or '9999', x['id']))
    write_table(destination, [row(x) for x in selected])
    print(f'已扫描 {args.school} 全部 {len(records)} 条保研信息；导出 {len(selected)} 条夏令营/预推免：{destination.resolve()}')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest='command', required=True)
    p = sub.add_parser('login', help='打开 Chrome 微信扫码，并保存登录状态')
    p.add_argument('--cdp', help='连接已开启远程调试的专用 Chrome，如 http://127.0.0.1:9229')
    p.add_argument('--timeout', type=int, default=600)
    p.set_defaults(func=login)
    p = sub.add_parser('export', help='导出该校网站现存全部夏令营/预推免信息')
    p.add_argument('--school', required=True, help='学校完整名称，如 北京大学')
    p.add_argument('--kind', choices=['全部', '夏令营', '预推免'], default='全部')
    p.add_argument('--year', type=int, help='按网站 year 字段精确筛选；默认不限制年份')
    p.add_argument('-o', '--output', required=True, help='输出 .xlsx / .csv / .json')
    p.add_argument('--anonymous', action='store_true', help='仅使用公开接口，不读取登录凭据')
    p.add_argument('--force', action='store_true', help='覆盖已有输出文件')
    p.add_argument('--page-size', type=int, choices=range(1, 101), default=100, metavar='1..100')
    p.add_argument('--delay', type=float, default=0.3, help='请求最小间隔秒数，默认 0.3')
    p.set_defaults(func=export)
    if 'report' in sys.argv[1:2]:
        from baoyan_report import add_arguments
        add_arguments(sub.add_parser('report', help='985/211预推免筛选，大学分组、活动结束降序并着色'))
    else:
        sub.add_parser('report', help='985/211预推免筛选，大学分组、活动结束降序并着色')
    args = parser.parse_args()
    try:
        args.func(args)
    except KeyboardInterrupt:
        print('已取消。', file=sys.stderr)
        return 130
    except (RuntimeError, httpx.HTTPError, ValueError, OSError, KeyError) as e:
        print(f'错误：{e}', file=sys.stderr)
        return 1
    return 0


if __name__ == '__main__':
    sys.exit(main())
