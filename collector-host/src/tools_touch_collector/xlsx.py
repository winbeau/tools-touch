from __future__ import annotations

import hashlib
import json
import os
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook


SCHEMA_VERSION = 1
MAX_CELL_LENGTH = 32_767
INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")
FORMULA_PREFIXES = ("=", "+", "-", "@")


def _safe_path(root: Path, relative: str) -> Path:
    path = Path(relative)
    if path.is_absolute() or any(part in ("", ".", "..") for part in path.parts):
        raise ValueError("XLSX_ROW_PATH_INVALID")
    resolved = (root / path).resolve()
    if not resolved.is_relative_to(root.resolve()):
        raise ValueError("XLSX_ROW_PATH_INVALID")
    return resolved


def _sheet_name(name: str, used: set[str], suffix: str = "") -> str:
    candidate = INVALID_SHEET_CHARS.sub("_", str(name)).strip().strip("'") or "Sheet"
    candidate = (candidate[: 31 - len(suffix)] + suffix)[:31]
    key = candidate.casefold()
    if key not in used:
        used.add(key)
        return candidate
    digest = hashlib.sha256((candidate + suffix).encode("utf-8")).hexdigest()[:6]
    candidate = (candidate[:24] + "~" + digest)[:31]
    counter = 2
    while candidate.casefold() in used:
        candidate = (candidate[:27] + f"~{counter}")[:31]
        counter += 1
    used.add(candidate.casefold())
    return candidate


def _cell(value: Any, long_text: Any, table: str, row_number: int, column: str) -> Any:
    if value is None or isinstance(value, (bool, int, float)):
        return value
    if not isinstance(value, str):
        value = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    if len(value) > MAX_CELL_LENGTH:
        text_id = hashlib.sha256(f"{table}\0{row_number}\0{column}".encode("utf-8")).hexdigest()[:16]
        for chunk_number, start in enumerate(range(0, len(value), MAX_CELL_LENGTH - 100), 1):
            long_text.append([text_id, table, row_number, column, chunk_number, value[start : start + MAX_CELL_LENGTH - 100]])
        return f"[long_text:{text_id}]"
    if value.startswith(FORMULA_PREFIXES):
        return "'" + value
    return value


def export_xlsx(request: dict[str, Any], output: Path) -> Path:
    if set(request) != {"schema_version", "tables", "max_rows_per_sheet"} or request["schema_version"] != SCHEMA_VERSION:
        raise ValueError("INVALID_XLSX_REQUEST")
    tables = request["tables"]
    max_rows = request["max_rows_per_sheet"]
    if not isinstance(tables, list) or not isinstance(max_rows, int) or not 1 <= max_rows <= 1_000_000:
        raise ValueError("INVALID_XLSX_REQUEST")
    output = output.resolve()
    output.mkdir(parents=True, exist_ok=True)
    root = output
    validated_tables: list[tuple[dict[str, Any], Path, list[str], int]] = []
    for table in tables:
        if not isinstance(table, dict) or set(table) != {"name", "path", "columns", "row_count"}:
            raise ValueError("INVALID_XLSX_TABLE")
        path = _safe_path(root, str(table["path"]))
        columns = table["columns"]
        expected_count = table["row_count"]
        if not path.is_file() or not isinstance(columns, list) or not all(isinstance(column, str) for column in columns) or not isinstance(expected_count, int):
            raise ValueError("INVALID_XLSX_TABLE")
        validated_tables.append((table, path, columns, expected_count))
    used: set[str] = set()
    workbook = Workbook(write_only=True)
    index = workbook.create_sheet(_sheet_name("Index", used))
    index.append(["table", "sheet", "row_count", "source_path"])
    long_text = None
    manifest_tables: list[dict[str, Any]] = []
    for table, path, columns, expected_count in validated_tables:
        name = str(table["name"])
        row_count = 0
        sheet_number = 1
        sheet = workbook.create_sheet(_sheet_name(name, used))
        sheet.append(columns)
        with path.open(encoding="utf-8") as stream:
            for row_count, line in enumerate(stream, 1):
                try:
                    row = json.loads(line)
                except json.JSONDecodeError as error:
                    raise ValueError("INVALID_XLSX_ROW") from error
                if not isinstance(row, dict):
                    raise ValueError("INVALID_XLSX_ROW")
                if row_count > 1 and (row_count - 1) % max_rows == 0:
                    sheet_number += 1
                    sheet = workbook.create_sheet(_sheet_name(name, used, f"~{sheet_number}"))
                    sheet.append(columns)
                if long_text is None:
                    long_text = workbook.create_sheet(_sheet_name("LongText", used))
                    long_text.append(["text_id", "table", "row_number", "column", "chunk_number", "text"])
                sheet.append([_cell(row.get(column), long_text, name, row_count, column) for column in columns])
        if row_count != expected_count:
            raise ValueError("XLSX_ROW_COUNT_MISMATCH")
        index.append([name, sheet.title if sheet_number == 1 else f"{sheet_number} sheets", row_count, str(table["path"])])
        manifest_tables.append({"name": name, "row_count": row_count, "sheet_count": sheet_number})

    output_file = output / "workbook.xlsx"
    temporary = output / ".workbook.xlsx.tmp"
    try:
        workbook.save(temporary)
        os.replace(temporary, output_file)
        digest = hashlib.sha256(output_file.read_bytes()).hexdigest()
        result = {"schema_version": SCHEMA_VERSION, "tables": manifest_tables,
                  "workbook": {"path": "workbook.xlsx", "sha256": digest, "byte_length": output_file.stat().st_size}}
        (output / "manifest.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        return output_file
    finally:
        temporary.unlink(missing_ok=True)


def validate_xlsx(output: Path) -> dict[str, Any]:
    manifest_path = output / "manifest.json"
    workbook_path = output / "workbook.xlsx"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if manifest.get("schema_version") != SCHEMA_VERSION or not workbook_path.is_file():
        raise ValueError("XLSX_MANIFEST_INVALID")
    if hashlib.sha256(workbook_path.read_bytes()).hexdigest() != manifest["workbook"]["sha256"]:
        raise ValueError("XLSX_HASH_MISMATCH")
    load_workbook(workbook_path, read_only=True, data_only=False).close()
    return manifest


def import_xlsx(request: dict[str, Any], output: Path) -> Path:
    if set(request) != {"schema_version", "source_path", "sheet_name"} or request["schema_version"] != SCHEMA_VERSION:
        raise ValueError("INVALID_XLSX_IMPORT_REQUEST")
    source = Path(str(request["source_path"])).expanduser()
    if not source.is_file() or source.is_symlink():
        raise ValueError("XLSX_SOURCE_INVALID")
    requested_sheet = request["sheet_name"]
    if requested_sheet is not None and not isinstance(requested_sheet, str):
        raise ValueError("XLSX_SHEET_INVALID")
    output = output.resolve()
    output.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(source, read_only=True, data_only=False)
    try:
        sheets = [sheet for sheet in workbook.worksheets if sheet.title.casefold() not in {"index", "longtext"}]
        if requested_sheet is not None:
            sheets = [sheet for sheet in sheets if sheet.title == requested_sheet]
        if len(sheets) != 1:
            raise ValueError("XLSX_SHEET_NOT_FOUND" if requested_sheet else "XLSX_SHEET_AMBIGUOUS")
        sheet = sheets[0]
        rows = sheet.iter_rows(values_only=True)
        try:
            raw_headers = next(rows)
        except StopIteration as error:
            raise ValueError("XLSX_HEADER_MISSING") from error
        columns = [str(value) if value is not None else "" for value in raw_headers]
        if not columns or any(not column.strip() for column in columns) or len({column.casefold() for column in columns}) != len(columns):
            raise ValueError("XLSX_HEADER_INVALID")
        rows_path = output / "rows.jsonl"
        row_count = 0
        with rows_path.open("w", encoding="utf-8", newline="\n") as stream:
            for row in rows:
                values = list(row)
                if not any(value is not None and str(value) != "" for value in values):
                    continue
                if len(values) < len(columns):
                    values.extend([None] * (len(columns) - len(values)))
                values = values[: len(columns)]
                stream.write(json.dumps([_import_cell(value) for value in values], ensure_ascii=False, separators=(",", ":")) + "\n")
                row_count += 1
        manifest = {"schema_version": SCHEMA_VERSION, "sheet_name": sheet.title, "columns": columns,
                    "row_count": row_count, "rows_path": "rows.jsonl"}
        (output / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        return output / "manifest.json"
    finally:
        workbook.close()


def _import_cell(value: Any) -> Any:
    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)
