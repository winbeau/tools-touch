from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from tools_touch_collector.xlsx import export_xlsx, import_xlsx, validate_xlsx


class XlsxExportTests(unittest.TestCase):
    def test_streamed_workbook_splits_sheets_and_preserves_formula_and_long_text(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            output = root / "xlsx"
            rows = output / "rows"
            rows.mkdir(parents=True)
            long_text = "长文本" * 8_000
            (rows / "one.jsonl").write_text(
                json.dumps({"id": "1", "text": "=SUM(A1:A2)"}, ensure_ascii=False) + "\n"
                + json.dumps({"id": "2", "text": long_text}, ensure_ascii=False) + "\n"
                + json.dumps({"id": "3", "text": "普通文本"}, ensure_ascii=False) + "\n",
                encoding="utf-8",
            )
            (rows / "two.jsonl").write_text(json.dumps({"id": "4", "text": "第二表"}) + "\n", encoding="utf-8")
            request = {
                "schema_version": 1,
                "tables": [
                    {"name": "A/B", "path": "rows/one.jsonl", "columns": ["id", "text"], "row_count": 3},
                    {"name": "A_B", "path": "rows/two.jsonl", "columns": ["id", "text"], "row_count": 1},
                ],
                "max_rows_per_sheet": 2,
            }
            workbook_path = export_xlsx(request, output)
            manifest = validate_xlsx(output)
            self.assertEqual(manifest["tables"][0]["row_count"], 3)
            workbook = load_workbook(workbook_path, read_only=True, data_only=False)
            self.assertGreaterEqual(len(workbook.sheetnames), 5)
            cells = [cell for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row]
            self.assertIn("'=SUM(A1:A2)", [cell.value for cell in cells])
            recovered = "".join(str(cell.value) for cell in cells if isinstance(cell.value, str) and cell.value.startswith("长文本"))
            self.assertEqual(recovered, long_text)
            workbook.close()

    def test_rejects_path_escape(self):
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaisesRegex(ValueError, "XLSX_ROW_PATH_INVALID"):
                export_xlsx({"schema_version": 1, "tables": [{"name": "x", "path": "../x", "columns": ["id"], "row_count": 0}], "max_rows_per_sheet": 10}, Path(directory))

    def test_imports_one_selected_sheet_as_json_rows(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source.xlsx"
            workbook = Workbook()
            first = workbook.active
            first.title = "Data"
            first.append(["key", "note"])
            first.append(["a", "=keep as text"])
            second = workbook.create_sheet("Other")
            second.append(["ignored"])
            workbook.save(source)
            workbook.close()
            output = root / "import"
            manifest_path = import_xlsx({"schema_version": 1, "source_path": str(source), "sheet_name": "Data"}, output)
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            self.assertEqual(manifest["columns"], ["key", "note"])
            self.assertEqual(manifest["row_count"], 1)
            self.assertEqual(json.loads((output / "rows.jsonl").read_text(encoding="utf-8").strip()), ["a", "=keep as text"])


if __name__ == "__main__":
    unittest.main()
